"""Local attachment discovery for offline/private benchmark and production handoff runs.

The ServiceNow ingestion/file-management layer is intentionally outside this module. It is expected
 to create one directory per CR and place that CR's downloaded attachments inside it. This module
consumes that prepared workspace without making ServiceNow calls or modifying the files.

Files remain on the caller's machine. Supported text/tabular documents are extracted locally. Image
attachments are retained as explicit unparsed evidence so the validator never silently ignores them.
"""
from __future__ import annotations

from pathlib import Path

from .evidence import EvidenceDocument

TEXT_SUFFIXES = {".pdf", ".xlsx", ".xlsm", ".txt", ".md", ".csv", ".eml", ".json"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | IMAGE_SUFFIXES


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install the 'docs' extra to parse PDF attachments") from exc
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install the 'docs' extra to parse XLSX attachments") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value not in (None, "")]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    if suffix in {".txt", ".md", ".csv", ".eml", ".json"}:
        return path.read_text(encoding="utf-8", errors="replace")
    return ""


def _cr_directory(root: Path, cr_number: str) -> Path | None:
    """Resolve the ingestion layer's canonical ``<root>/<CR number>`` directory.

    The exact directory match is preferred so an attachment belonging to CR123 is never selected
    merely because another path happens to contain the same string. A recursive fallback is kept
    for compatibility with the earlier private benchmark layout where CR folders could be nested.
    """
    needle = cr_number.strip()
    if not needle:
        return None

    direct = root / needle
    if direct.is_dir():
        return direct

    for candidate in root.rglob(needle):
        if candidate.is_dir() and candidate.name == needle:
            return candidate
    return None


def _root_attachments(root: Path, cr_number: str) -> list[Path]:
    """Return legacy root-level files explicitly named for this CR.

    The production contract is one folder per CR, but older callers/tests may place files directly
    under the attachment root. Only names with the exact CR number followed by a separator are
    accepted; unrelated sibling CR evidence is never guessed.
    """
    needle = cr_number.strip()
    if not needle:
        return []
    matches: list[Path] = []
    for path in root.iterdir():
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        stem = path.stem
        if stem == needle or stem.startswith(f"{needle}_") or stem.startswith(f"{needle}-"):
            matches.append(path)
    return sorted(matches)


def discover_attachments(root: Path, cr_number: str) -> list[Path]:
    """Find supported files in the CR-specific workspace directory.

    The ingestion layer owns folder creation and file placement. This function only reads files
    from the resolved CR directory and never guesses unrelated evidence from sibling CR folders.
    For backward compatibility, explicitly CR-prefixed files directly under ``root`` are also
    accepted when no canonical CR directory exists.
    """
    if not root.exists() or not root.is_dir():
        return []

    cr_dir = _cr_directory(root, cr_number)
    if cr_dir is not None:
        return sorted(
            path
            for path in cr_dir.rglob("*")
            if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
        )

    return _root_attachments(root, cr_number)


def load_attachments_for_cr(root: Path, cr_number: str) -> list[EvidenceDocument]:
    documents: list[EvidenceDocument] = []
    for path in discover_attachments(root, cr_number):
        suffix = path.suffix.lower()
        metadata = {
            "local_path": str(path),
            "bytes": path.stat().st_size,
            "requires_vision": suffix in IMAGE_SUFFIXES,
            "extraction_error": None,
        }
        if suffix in IMAGE_SUFFIXES:
            text = ""
        else:
            try:
                text = extract_text(path)
                if not text.strip():
                    metadata["extraction_error"] = "No extractable text found; document may be scanned/image-only or empty."
            except (OSError, RuntimeError, ValueError) as exc:
                text = ""
                metadata["extraction_error"] = f"{type(exc).__name__}: {exc}"
        documents.append(
            EvidenceDocument(
                ref=str(path),
                name=path.name,
                text=text,
                document_type=suffix.lstrip(".") or "unknown",
                metadata=metadata,
            )
        )
    return documents
