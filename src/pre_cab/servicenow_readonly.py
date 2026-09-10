"""Read-only ServiceNow REST client for controlled ingestion.

The client exposes GET-only Change Request and attachment operations. Credentials are never included in
returned metadata or logs. Binary attachments can be downloaded in-memory and converted into
EvidenceDocument records without granting any write capability.
"""
from __future__ import annotations

import base64
import io
import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

from .evidence import EvidenceDocument


@dataclass(frozen=True)
class ServiceNowCredentials:
    base_url: str
    username: str
    password: str


class ReadOnlyServiceNowClient:
    """Minimal, bounded, GET-only ServiceNow client.

    Retries are limited to transient transport failures and 429/5xx responses. Authentication and
    other 4xx responses fail immediately. Attachment downloads are size-bounded to avoid untrusted
    ServiceNow content exhausting process memory.
    """

    def __init__(
        self,
        credentials: ServiceNowCredentials,
        table: str = "change_request",
        *,
        timeout_seconds: float = 30.0,
        attachment_timeout_seconds: float = 60.0,
        max_retries: int = 2,
        max_attachment_bytes: int = 25 * 1024 * 1024,
        retry_backoff_seconds: float = 0.5,
    ) -> None:
        if max_retries < 0:
            raise ValueError("max_retries must be >= 0")
        if max_attachment_bytes <= 0:
            raise ValueError("max_attachment_bytes must be > 0")
        self.credentials = credentials
        self.table = table
        self.timeout_seconds = timeout_seconds
        self.attachment_timeout_seconds = attachment_timeout_seconds
        self.max_retries = max_retries
        self.max_attachment_bytes = max_attachment_bytes
        self.retry_backoff_seconds = retry_backoff_seconds

    def _request(self, path: str, *, accept: str = "application/json") -> Request:
        request = Request(
            self.credentials.base_url.rstrip("/") + "/" + path.lstrip("/"),
            method="GET",
        )
        token = base64.b64encode(
            f"{self.credentials.username}:{self.credentials.password}".encode()
        ).decode()
        request.add_header("Authorization", f"Basic {token}")
        request.add_header("Accept", accept)
        request.add_header("User-Agent", "pre-cab-validator/read-only")
        return request

    def _get(self, path: str) -> dict[str, Any]:
        raw = self._read(path, accept="application/json", timeout=self.timeout_seconds)
        data = json.loads(raw.decode("utf-8"))
        return data if isinstance(data, dict) else {"result": data}

    def _get_bytes(self, path: str) -> bytes:
        return self._read(path, accept="*/*", timeout=self.attachment_timeout_seconds)

    def _read(self, path: str, *, accept: str, timeout: float) -> bytes:
        last_error: Exception | None = None
        for attempt in range(self.max_retries + 1):
            try:
                with urlopen(self._request(path, accept=accept), timeout=timeout) as response:
                    content_length = response.headers.get("Content-Length")
                    if content_length and int(content_length) > self.max_attachment_bytes:
                        raise ValueError("ServiceNow attachment exceeds configured size limit")
                    chunks: list[bytes] = []
                    total = 0
                    while True:
                        chunk = response.read(min(1024 * 1024, self.max_attachment_bytes - total + 1))
                        if not chunk:
                            break
                        total += len(chunk)
                        if total > self.max_attachment_bytes:
                            raise ValueError("ServiceNow response exceeds configured size limit")
                        chunks.append(chunk)
                    return b"".join(chunks)
            except HTTPError as exc:
                last_error = exc
                if exc.code not in {429, 500, 502, 503, 504} or attempt >= self.max_retries:
                    raise
            except (URLError, TimeoutError, OSError) as exc:
                last_error = exc
                if attempt >= self.max_retries:
                    raise
            if self.retry_backoff_seconds > 0:
                time.sleep(self.retry_backoff_seconds * (2**attempt))
        raise RuntimeError("ServiceNow request failed") from last_error

    def get_change(self, number: str) -> dict[str, Any]:
        encoded = quote(f"number={number}", safe="")
        data = self._get(
            f"api/now/table/{self.table}?sysparm_query={encoded}&sysparm_limit=1&sysparm_display_value=true"
        )
        rows = data.get("result") or []
        if not rows:
            raise KeyError(f"Change request not found: {number}")
        return dict(rows[0])

    def list_attachments(self, sys_id: str) -> list[dict[str, Any]]:
        encoded = quote(f"table_sys_id={sys_id}", safe="")
        data = self._get(f"api/now/attachment?sysparm_query={encoded}")
        return [dict(row) for row in data.get("result") or []]

    def download_attachment(self, attachment_sys_id: str) -> bytes:
        return self._get_bytes(f"api/now/attachment/{quote(attachment_sys_id, safe='')}/file")

    @staticmethod
    def _extract_attachment_text(name: str, content: bytes) -> str:
        suffix = Path(name).suffix.lower()
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError("Install the 'docs' extra to parse PDF attachments") from exc
            reader = PdfReader(io.BytesIO(content))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("Install the 'docs' extra to parse XLSX attachments") from exc
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in workbook.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value not in (None, "")]
                    if values:
                        lines.append(" | ".join(values))
            return "\n".join(lines)
        return content.decode("utf-8", errors="replace")

    def evidence_documents(self, change: dict[str, Any]) -> list[EvidenceDocument]:
        sys_id = str(change.get("sys_id") or "").strip()
        if not sys_id:
            return []
        documents: list[EvidenceDocument] = []
        for attachment in self.list_attachments(sys_id):
            attachment_id = str(attachment.get("sys_id") or "").strip()
            if not attachment_id:
                continue
            name = str(attachment.get("file_name") or attachment_id)
            try:
                content = self.download_attachment(attachment_id)
                text = self._extract_attachment_text(name, content)
                extraction_error = None
            except (OSError, RuntimeError, ValueError, HTTPError, URLError) as exc:
                text = ""
                extraction_error = f"{type(exc).__name__}: {exc}"
            documents.append(
                EvidenceDocument(
                    ref=f"servicenow:{attachment_id}",
                    name=name,
                    text=text,
                    document_type=str(attachment.get("content_type") or Path(name).suffix.lstrip(".") or "unknown"),
                    metadata={
                        "attachment_sys_id": attachment_id,
                        "size_bytes": attachment.get("size_bytes"),
                        "content_type": attachment.get("content_type"),
                        "extraction_error": extraction_error,
                    },
                )
            )
        return documents

    def ingest_change(self, number: str) -> tuple[dict[str, Any], list[EvidenceDocument]]:
        change = self.get_change(number)
        return change, self.evidence_documents(change)
