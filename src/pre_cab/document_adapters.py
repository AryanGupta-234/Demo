"""Optional document extraction adapters.

Stage 2 operates on :class:`EvidenceDocument`; these adapters keep the core verifier independent
from PDF/XLSX/email parsing libraries. Optional dependencies are imported only when used.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from .evidence import EvidenceDocument


def load_text_file(path: str | Path, *, metadata: dict[str, Any] | None = None) -> EvidenceDocument:
    p = Path(path)
    return EvidenceDocument(ref=str(p), name=p.name, text=p.read_text(encoding="utf-8", errors="replace"), document_type=p.suffix.lstrip("."), metadata=metadata or {})


def load_pdf(path: str | Path, *, metadata: dict[str, Any] | None = None) -> EvidenceDocument:
    """Extract searchable text from a PDF using pypdf when installed."""
    p = Path(path)
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Install 'pypdf' to read PDF attachments") from exc
    reader = PdfReader(str(p))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)
    return EvidenceDocument(ref=str(p), name=p.name, text=text, document_type="pdf", metadata=metadata or {})


def load_xlsx(path: str | Path, *, metadata: dict[str, Any] | None = None) -> EvidenceDocument:
    """Flatten XLSX cells into searchable evidence text using openpyxl."""
    p = Path(path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install 'openpyxl' to read XLSX attachments") from exc
    workbook = load_workbook(filename=p, read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        chunks.append(f"[SHEET {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                chunks.append(" | ".join(values))
    return EvidenceDocument(ref=str(p), name=p.name, text="\n".join(chunks), document_type="xlsx", metadata=metadata or {})
